# -*- coding: utf-8 -*-
"""Build an Excel summary of all gacha/draw mechanics in the psw mod
(星之魔女祈愿). Output goes to mod_output/ which is git-ignored.

Sources:
- E:\\tmp\\psw_decoded\\loc\\*.tsv   (localisation key->text)
- E:\\tmp\\psw_decoded\\db\\**\\*.tsv (decoded DB tables)
- E:\\tmp\\psw_decoded\\psw_dilemma_report.tsv (pre-joined dilemma report)
- E:\\tmp\\psw_extract\\script\\campaign\\mod\\*.lua (pool definitions)
"""
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

DEC = Path(r"E:\tmp\psw_decoded")
LUA = Path(r"E:\tmp\psw_extract\script\campaign\mod")
OUT = Path(r"E:\Code\WH3-modding\mod_output\psw_魔女mod抽卡与效果汇总.xlsx")

# ---------------------------------------------------------------- helpers

def load_loc():
    loc = {}
    for f in (DEC / "loc").glob("*.tsv"):
        for i, line in enumerate(f.read_text(encoding="utf-8-sig").splitlines()):
            if i == 0 or line.startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) >= 2 and parts[0]:
                loc[parts[0]] = parts[1]
    # vanilla CN localisation for wh_main_effect_* etc.
    cn = Path(r"E:\Code\WH3-modding\bilingual\localisation__.loc_CN.tsv")
    if cn.exists():
        for i, line in enumerate(cn.read_text(encoding="utf-8-sig").splitlines()):
            if i == 0:
                continue
            parts = line.split("\t")
            if len(parts) >= 2 and parts[0] and parts[0] not in loc:
                loc[parts[0]] = parts[1]
    return loc

TAG_RE = re.compile(r"\[\[/?(?:col|img)(?::[^\]]*)?\]\]")

def clean(s: str) -> str:
    if not s:
        return ""
    s = TAG_RE.sub("", s)
    return s.replace("%%", "%").replace("\\n", "\n").strip()

def read_tsv(path: Path):
    """-> (fields, rows[dict])"""
    lines = [l for l in path.read_text(encoding="utf-8-sig").splitlines()
             if l and not l.startswith("#")]
    if not lines:
        return [], []
    fields = lines[0].split("\t")
    rows = []
    for line in lines[1:]:
        vals = line.split("\t")
        rows.append({f: (vals[i] if i < len(vals) else "") for i, f in enumerate(fields)})
    return fields, rows

def read_db_dir(name: str):
    rows = []
    for f in (DEC / "db" / name).glob("*.tsv"):
        rows += read_tsv(f)[1]
    return rows

def lua_pool_keys(path: Path, start_pat: str, end_pat=r"\n\}"):
    """Extract quoted keys from a `local X = { ... }` literal block."""
    text = path.read_text(encoding="utf-8")
    m = re.search(start_pat, text)
    if not m:
        return []
    tail = text[m.end():]
    end = re.search(end_pat, tail)
    block = tail[: end.start()] if end else tail
    return re.findall(r'"([^"]+)"', block)

def func_body(text: str, fname: str) -> str:
    m = re.search(r"function\s+" + re.escape(fname) + r"\b", text)
    if not m:
        return ""
    nxt = re.search(r"\nfunction\s", text[m.end():])
    return text[m.end(): m.end() + nxt.start()] if nxt else text[m.end():]

# ---------------------------------------------------------------- load

loc = load_loc()

def name_of(key, *prefixes):
    for p in prefixes or ("",):
        t = loc.get(p + key)
        if t:
            return clean(t)
    return ""

def desc_of(key):
    for p in ("effects_description_", "", "_effect"):
        t = loc.get(p + key) or loc.get(key + p if p.startswith("_") else p + key)
        if t:
            return clean(t)
    return ""

wb = openpyxl.Workbook()
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

def make_sheet(title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill = HDR_FILL; c.font = HDR_FONT
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"
    return ws

# =================================================================
# Sheet 0: 概览
# =================================================================
overview = [
    ["记忆结晶抽卡", "psw_jyjj_main_pickup / 各入口", "5 层稀有度等权 20%，层内均匀；sx=属性卡 xg=效果卡 xd=行动卡", "回合奖励/轮次奖励/卡牌效果触发；旅途商店购买指定卡", "无保底；未解锁卡不进入池子；手牌上限9，超出弃置"],
    ["祈愿天赋抽选", "psw_main_pickup（派系按钮）", "每位独立判定：幸运值→唯一天赋概率（上限25%），否则普通天赋", "消耗 1 点 psw_main_gmsj；每次抽出 jishu_xuanze 个选项（默认3上限4）", "免费重抽 free_repick_num；付费重抽 repick_cost 逐次+1；终端天赋按属性阈值保底"],
    ["祈愿装备（唯一物品抽奖）", "psw_main_pickup 唯一物品分支", "按 cj_* 属性区间命中，命中率=各档位概率之和", "消耗 10 点 psw_main_qysj；nor_count>8 且 qysj_cj≥10 且回合≥20（或 25回合后选择祈愿装备）", "各唯一物品全局仅一次，命中后从池移除；帕秋莉决定觉醒/未觉醒版"],
    ["游戏王抽卡", "psw_main_yqdjb_kachi", "random(1,100) 固定区间，见明细表", "拥有天赋『友情的羁绊』(psw_tf_nor_yqdjb_nor1) 后一次性触发", "一次性；集换有额外组合奖励（艾克佐迪亚/龙骑士盖亚/超魔导剑士/三幻神）"],
    ["抽取学生（基沃托斯）", "psw_dlc7_zhaomu_stu_effect", "所有未解锁学生均匀随机", "消耗 25 点 psw_main_qysj，一次抽 pick_num 个（默认3上限5）", "已解锁学生不再重复出现"],
    ["扭蛋（角色事件）", "dilemma_tf_*（73个扭蛋事件）", "事件二选一：获得对应天赋/效果组 或『不想丸辣』放弃", "事件弹窗触发", "效果为 1 回合 effect_bundle，详见『扭蛋与祈愿事件』表"],
    ["救济的祈愿", "set_psw_tf_ex_yhzl_2 觉醒", "非随机", "鹿目圆『渺小的希望』+因果≥10 自动升级", "属性奖励：治愈+2、神性+3、免死+1"],
    ["菲比祈愿（遥远的愿景）", "dilemma_tf_dlc7_feibi_qiyuan", "非随机，8选1", "feibi_qiyuan 计数每满30触发", "七维属性各+1 或全体回复3HP"],
    ["星星祈愿（不死者之王）", "dilemma_tf_zsqj_xuyuan", "非随机，5选1", "点击不死者之王建筑，消耗10 bszzw_mlcl", "神性/幸运/魔力/剑术/魅力 +2"],
    ["菲娅的祝福", "dilemma_tf_fia_fydzf_1/2", "非随机，分页选项", "计数消耗型", "祈愿保底（下次必唯一/专属）、唯一物品、异变切换、属性+1等"],
    ["引发奇迹", "dilemma_tf_dlc7_miao_yfqj", "非随机，9选1", "事件触发", "九个奇迹选项，详见明细表"],
    ["移岁诛邪", "dilemma_tf_dlc7_jinxi_yszh", "非随机，5选1", "事件触发", "记忆结晶数+5/共鸣水晶+3/下次祈愿必唯一/必专属/神性幸运+1"],
    ["守矢神社", "dilemma_tf_dlc7_ssss", "非随机，3选1", "事件触发", "取出奉纳/升级塞钱箱/取消"],
    ["预知梦（记忆结晶）", "dilemma_tf_jyjj_yuzhi", "选类型后从稀有度5抽1", "消耗1张『预知梦』", "sx/xg/xd 三选一；xd 额外加入飞雷神"],
    ["遗忘之祭仪", "psw_ywzjy_main_effect", "稀有度5已解锁中随机", "消耗1张『秘术·遗忘之祭仪』", "幸运≥10 时再抽1张"],
]
make_sheet("机制概览", ["机制", "内部入口", "随机方式/概率", "消耗与触发", "保底/特殊规则"], overview,
           [22, 34, 44, 40, 50])

# =================================================================
# Sheet 1: 扭蛋与祈愿事件（预 join 的 TSV）
# =================================================================
fields, rep = read_tsv(DEC / "psw_dilemma_report.tsv")
rows = []
for r in rep:
    rows.append([
        r["dilemma_key"], clean(r["event_name"]), r["choice_key"],
        clean(r["choice_text"]), r["payload_type"], r["payload_target"],
        r["value"], r["duration"], r["effect_key"], r["effect_scope"],
        clean(r["effect_text"]),
    ])
make_sheet("扭蛋与祈愿事件", fields, rows, [34, 14, 10, 30, 14, 34, 8, 8, 40, 26, 60])

# =================================================================
# Sheet 2: 记忆结晶卡池
# =================================================================
jyjj_src = (LUA / "psw_记忆结晶.lua").read_text(encoding="utf-8")
CARD_EFFECT = {
 "psw_jyjj_sx_main_fangyu_1": "随机1-3点临时防御，1回合，可叠加",
 "psw_jyjj_sx_main_gongji_1": "随机1-3点临时攻击，1回合，可叠加",
 "psw_jyjj_xg_main_moli_1": "国库 += 100×回合数 + 500×魔力值",
 "psw_jyjj_xg_main_zhiyu_1": "全军回复10% HP，额外+2临时HP",
 "psw_jyjj_xg_main_heijian": "+1临时攻击 +1临时防御，1回合",
 "psw_jyjj_sx_jinxi_yishi_1": "调用 psw_jyjj_effect_jinxi_fgjh_1（今汐·浮光霁寒一式）",
 "psw_jyjj_sx_yozi_cqbns": "临时攻击+3，1回合",
 "psw_jyjj_sx_main_fangyu_2": "随机1-5点临时防御，1回合",
 "psw_jyjj_sx_main_gongji_2": "随机1-5点临时攻击，1回合",
 "psw_jyjj_xg_main_moli_2": "国库 += 200×回合数 + 1000×魔力值",
 "psw_jyjj_xg_main_zhiyu_2": "全军回复20% HP，额外+3临时HP",
 "psw_jyjj_xg_main_xushu": "随机1-3点临时攻击+临时防御",
 "psw_jyjj_sx_jinxi_yishi_2": "调用 psw_jyjj_effect_jinxi_fgjh_2",
 "psw_jyjj_sx_jinxi_yishi_3": "调用 psw_jyjj_effect_jinxi_fgjh_3",
 "psw_jyjj_sx_yozi_kslbkm": "魔王-3HP，临时攻击+5，1回合",
 "psw_jyjj_xd_feibi_zhudao": "调用 psw_dlc7_jyjj_zhudao_effect",
 "psw_jyjj_xd_mingren_lxw": "代码已注释，暂无明显效果",
 "psw_jyjj_xd_xingye_sqss": "apply_num=3（手枪速射）",
 "psw_jyjj_xd_zhu_guangyu": "调用 psw_dlc7_jyjj_effect_guangyu",
 "psw_jyjj_sx_main_fangyu_3": "随机1-7点临时防御，1回合",
 "psw_jyjj_sx_main_gongji_3": "随机1-7点临时攻击，1回合",
 "psw_jyjj_xg_main_moli_3": "国库 += 400×回合数 + 2000×魔力值",
 "psw_jyjj_xg_main_zhiyu_3": "全军回复30% HP，额外+5临时HP",
 "psw_jyjj_xg_main_mlbf": "apply_num=魔力涌动层数",
 "psw_jyjj_xg_main_clzy": "全人类回复3点HP",
 "psw_jyjj_xg_main_mhpy": "调用 psw_tfmw_touzi_random（随机骰子）",
 "psw_jyjj_xd_main_jrdf": "使用次数+1（回复）",
 "psw_jyjj_xg_main_nvpu": "临时防御+6，2回合，不可叠加",
 "psw_jyjj_xg_main_poyin": "随机1-5点临时攻击+临时防御",
 "psw_jyjj_sx_jinxi_yishi_4": "调用 psw_jyjj_effect_jinxi_fgjh_4（终式）",
 "psw_jyjj_sx_klt_blmx": "调用 psw_dlc7_jyjj_klt_blmx",
 "psw_jyjj_sx_tsg_yy": "调用 psw_dlc7_tsg_jyjj_effect_yy",
 "psw_jyjj_sx_xmy_xlhxcd": "调用 psw_dlc7_jyjj_xmy_xlhxcd_effect",
 "psw_jyjj_sx_yozi_gzywl": "魔王-3HP，临时攻击+7，1回合",
 "psw_jyjj_xd_feibi_jzh": "调用 psw_dlc7_jyjj_effect_jzh（目标角色）",
 "psw_jyjj_xd_klt_sizhao": "apply_num=3；珂莱塔鸣奏时×1.5；调用 psw_dlc7_jyjj_klt_sizhao",
 "psw_jyjj_xd_yly_smck": "目标型仪式（使魔操控）",
 "psw_jyjj_xd_zuozhu_jjtm": "调用 psw_jyjj_jjtm_effect",
 "psw_jyjj_xg_feibi_zxgmzs": "调用 psw_dlc7_jyjj_effect_feibi_zxgmzs",
 "psw_jyjj_xg_firefly_fyzh": "调用 psw_dlc7_firefly_fyzh_jyjj_effect",
 "psw_jyjj_xg_lilika_nshf": "调用 psw_dlc7_jyjj_lilika_nshf_effect",
 "psw_jyjj_xg_qiesi_qyd": "目标型仪式，psw_dlc7_qeisi_qyd_effect",
 "psw_jyjj_xg_xiaochun_sqzz": "魔王-3HP，临时攻击+3，3回合",
 "psw_jyjj_xg_xnm_ycks": "目标型仪式，psw_dlc7_xnm_ycks_effect",
 "psw_jyjj_xg_zaomiao_hf": "调用 psw_dlc7_jyjj_miao_hf_effect",
 "psw_jyjj_xg_zuozhu_tstz": "目标型仪式，psw_jyjj_tstz_effect",
 "psw_jyjj_xg_main_jianpiao": "目标型仪式（该剪票了哦）",
 "psw_jyjj_xg_shaonv_chuhui": "目标型仪式（月兆祝赐·初辉）",
 "psw_jyjj_xg_xinghui_shxm": "调用 psw_dlc7_xinghui_jyjj_shxm",
 "psw_jyjj_sx_main_fangyu_4": "随机1-10点临时防御，1回合",
 "psw_jyjj_sx_main_gongji_4": "随机1-10点临时攻击，1回合",
 "psw_jyjj_xg_main_moli_4": "国库 += 500×回合数 + 2500×魔力值",
 "psw_jyjj_xg_main_zhiyu_4": "全军回复50% HP，额外+8临时HP",
 "psw_jyjj_xd_main_tyms": "立即再抽2张记忆结晶",
 "psw_jyjj_xg_main_cisi": "魔王-6HP；临时攻击+5，3回合",
 "psw_jyjj_xg_main_caibao": "共鸣水晶(gmsj)+1",
 "psw_jyjj_xg_main_chelun": "随机1-8点临时攻击+临时防御",
 "psw_jyjj_sx_klt_swcc": "调用 psw_dlc7_jyjj_klt_swcc",
 "psw_jyjj_xd_anne_jdyl": "目标型仪式，psw_jyjj_effect_anne_jdyl",
 "psw_jyjj_xd_mingren_xflxw": "代码已注释，暂无明显效果",
 "psw_jyjj_xd_weilinai_xxhzf": "调用 psw_dlc7_jyjj_xxhzf_wln_effect",
 "psw_jyjj_xd_xiaoye_shifu": "目标型仪式（咲夜特制秒表）",
 "psw_jyjj_xd_yozi_qxwh": "直接获得『攻击(神刀)』sx gongji_5；施加 psw_jyjj_qxwh_effect",
 "psw_jyjj_xd_zaomiao_ql": "调用 psw_dlc7_jyjj_miao_ql_effect",
 "psw_jyjj_xd_zaomiao_xl": "apply_num=魔力涌动；施加『星落』效果组2回合",
 "psw_jyjj_xd_zaomiao_ywzjy": "触发遗忘之祭仪：稀有度5已解锁中随机抽1（幸运≥10再抽1）",
 "psw_jyjj_xg_aly_yjyh": "调用 psw_dlc7_jyjj_effect_yjyh",
 "psw_jyjj_xg_anne_gyzmd": "调用 psw_jyjj_effect_anne_gyzmd",
 "psw_jyjj_xg_anne_zybj": "apply_num=floor(魔力值×1.5)",
 "psw_jyjj_xg_bcjh_bqtc": "外部函数/仪式处理",
 "psw_jyjj_xg_feibi_guanghui": "按菲比贯星咏调伤害计算；否则调用 psw_dlc7_jyjj_feibi_gh_effect",
 "psw_jyjj_xg_jinxi_snfm": "调用 psw_jyjj_effect_jinxi_fgjh_snfm；apply_num=5",
 "psw_jyjj_xg_klt_zsyz": "apply_num=10；珂莱塔鸣奏×1.5；珂莱塔+2协奏能量",
 "psw_jyjj_xg_ktll_gzjq": "目标型仪式（感知汲取）",
 "psw_jyjj_xg_mali_chu": "外部函数/仪式处理（慈爱的飞吻）",
 "psw_jyjj_xg_xingye_yzgj": "压制攻击：清空手牌并按手牌数重新抽取，额外+3张手枪速射",
 "psw_jyjj_xg_ams_ghzy": "调用 psw_dlc7_ams_jyjj_ghzy_effect",
 "psw_jyjj_xg_xlfh_jhly": "调用 psw_dlc7_jyjj_xlfh_jhly_effect",
 "psw_jyjj_xd_delisha_gyzy": "调用 psw_dlc7_jyjj_gyzy_effect",
 "psw_jyjj_sx_main_fangyu_5": "随机1-20点临时防御，1回合",
 "psw_jyjj_sx_main_gongji_5": "随机1-20点临时攻击，1回合",
 "psw_jyjj_xd_main_yuzhi": "触发预知梦三选一（sx/xg/xd），从稀有度5抽1",
 "psw_jyjj_xg_main_zhongli": "坚固层数+10，结界层数+1",
 "psw_jyjj_sx_firefly_rsdj": "调用 psw_dlc7_jyjj_effect_firefly_rsdj + 播放电影",
 "psw_jyjj_sx_jplr_jydsjh": "外部函数/仪式处理（久远第四加护）",
 "psw_jyjj_xd_kokoro_jlzdy": "调用 psw_dlc7_emt_jyjj_jlzyd_effect（精灵之引导）",
 "psw_jyjj_xd_xmy_rpg": "调用 psw_dlc7_jyjj_xmy_rpg_effect",
 "psw_jyjj_xd_zaomiao_hsfkzr": "apply_num=floor(魔力涌动×2+幸运)；再动+3",
 "psw_jyjj_xg_ams_smys": "调用 psw_dlc7_ams_jyjj_smys_effect",
 "psw_jyjj_xg_ams_yjxs": "魔力计数+1；apply_num=floor(魔力涌动×3)",
 "psw_jyjj_xg_feixiao_gwzz": "调用 psw_dlc7_jyjj_fx_gwzz_effect",
 "psw_jyjj_xg_feixiao_kjfl": "调用 psw_dlc7_jyjj_fx_kjfl_effect",
 "psw_jyjj_xg_feixiao_xsqj": "调用 psw_dlc7_jyjj_fx_xsqj_effect",
 "psw_jyjj_xg_jinxi_jlpk": "调用 psw_jyjj_effect_jinxi_fgjh_jlpk；apply_num=10（受今汐韶光加成）",
 "psw_jyjj_xg_lilika_zhufu": "调用 psw_dlc7_jyjj_lilika_qynszzf_effect",
 "psw_jyjj_xg_yozi_wmdmx": "调用 psw_dlc7_wmdmx_jyjj_effect",
 "psw_jyjj_xg_zaomiao_smg": "调用 psw_jyjj_effect_miao_smg",
 "psw_jyjj_xg_sar_hdll": "调用 psw_dlc7_jyjj_sar_hdll",
 "psw_jyjj_xg_main_feileishen": "传送派系领袖到目标角色附近（使用类型判定为xd）",
 "psw_jyjj_xg_feilin_tuxi": "调用 psw_dlc7_feilin_jyjj_tuxi_effect",
 "psw_jyjj_xg_xlfh_qishi": "龙族计数+1，播放专属音乐",
}
TYPE_NAME = {"sx": "属性卡(耗体力)", "xg": "效果卡(耗次数)", "xd": "行动卡"}
jyjj_rows = []
for lv in range(1, 6):
    keys = lua_pool_keys(LUA / "psw_记忆结晶.lua", rf"local jyjj_pool_lv{lv} = \{{")
    for k in keys:
        t = re.search(r"psw_jyjj_(sx|xg|xd)_", k)
        jyjj_rows.append([
            f"稀有度{lv}（层级权重20%，层内1/{len(keys)}）",
            k, TYPE_NAME.get(t.group(1) if t else "", "?"),
            name_of(k, "rituals_display_name_"),
            f"20% ÷ {len(keys)} ≈ {20/len(keys):.2f}%（默认池全解锁）",
            CARD_EFFECT.get(k, ""),
        ])
make_sheet("记忆结晶卡池", ["稀有度/卡池", "卡牌key", "类型", "显示名", "单卡概率", "使用效果"], jyjj_rows,
           [30, 38, 16, 22, 24, 62])

# =================================================================
# Sheet 3-4: 天赋池（普通 / 唯一）
# =================================================================
yw_src = (LUA / "psw_额外函数.lua").read_text(encoding="utf-8")

def talent_desc(key):
    """Talent -> set bundle -> effects -> loc descriptions."""
    out = []
    for bk in (f"set_{key}", f"set_main_{key}", f"dilemma_tf_{key}"):
        for j in bundle_junctions.get(bk, []):
            d = loc.get("effects_description_" + j["effect_key"], "")
            if d:
                out.append(clean(d))
    if not out:
        # fallback: any loc effects_description_ ending with key
        for lk, lt in loc.items():
            if lk.startswith("effects_description_") and lk.endswith(key):
                out.append(clean(lt))
                break
    return "\n".join(dict.fromkeys(out))

bundle_junctions = {}
for j in read_db_dir("effect_bundles_to_effects_junctions_tables"):
    bundle_junctions.setdefault(j.get("effect_bundle_key", ""), []).append(j)

nor_body = func_body(yw_src, "psw_main_nor_pool")
nor_base = re.findall(r'"([^"]+)"', nor_body.split("local tf_pool = {", 1)[1].split("}", 1)[0]) if "local tf_pool = {" in nor_body else []
nor_extra = re.findall(r'table\.insert\([^,]+,\s*"([^"]+)"\)', nor_body.split("}", 1)[-1])
nor_rows = [[k, name_of(k, "building_culture_variants_name_"), "基础池", talent_desc(k)] for k in dict.fromkeys(nor_base)]
nor_rows += [[k, name_of(k, "building_culture_variants_name_"), "条件加入", talent_desc(k)] for k in dict.fromkeys(nor_extra) if k not in nor_base]
make_sheet("天赋池-普通", ["天赋key", "显示名", "加入方式", "效果说明"], nor_rows, [38, 26, 12, 80])

ex_base = lua_pool_keys(LUA / "psw_额外函数.lua", r"local ex_tf_to_insert_pool = \{")
ex_body = func_body(yw_src, "psw_main_ex_pool")
ex_extra = re.findall(r'(?:table\.insert\([^,]+,\s*|insert_more_tf_to_pool\()\s*"([^"]+)"', ex_body)
sp_rows = [[k, name_of(k, "building_culture_variants_name_"), "基础EX池", talent_desc(k)] for k in dict.fromkeys(ex_base)]
sp_rows += [[k, name_of(k, "building_culture_variants_name_"), "条件加入/加权", talent_desc(k)] for k in dict.fromkeys(ex_extra) if k not in ex_base]
make_sheet("天赋池-唯一EX-SP", ["天赋key", "显示名", "加入方式", "效果说明"], sp_rows, [38, 26, 14, 80])

terminal = [
    ["幸运≥10", "psw_tf_sp_xxdqj_ex", "星星与奇迹的魔女"],
    ["魔力≥10", "psw_tf_ex_zzhgt_ex", "最终号个体"],
    ["知性≥10", "psw_tf_ex_pql", "知识与避世的少女"],
    ["魅力≥10", "psw_tf_ex_shaonv", "新月诞生的少女"],
    ["魔法少女计数≥4", "psw_tf_ex_xzwlsy", "向着未来闪耀"],
    ["鹿目圆因果≥10", "psw_tf_ex_yhzl_1", "渺小的希望"],
    ["白夜100计数", "psw_tf_sp_baiye", "魔法少女与恶曾经敌对"],
    ["蒙面泳装团", "psw_tf_nor_mmyzt", "集结！蒙面泳装团！"],
]
make_sheet("终端天赋保底", ["触发条件", "天赋key", "显示名", "效果说明"],
           [[c, k, n, talent_desc(k)] for c, k, n in terminal], [24, 30, 26, 80])

# =================================================================
# Sheet 5: 游戏王抽卡
# =================================================================
ygo = [
 ["1-8","jnz_tf_nor_yqdjb_1_szss","死者苏生","8%","领袖可用「魔雨雾」次数+1，治疗次数+1"],
 ["9-16","jnz_tf_nor_yqdjb_1_gzhfj","光之护封剑","8%","领袖伤害减免+8%"],
 ["17-21","jnz_tf_nor_yqdjb_1_hmss","黑魔术师","5%","所有法术法力消耗-10%"],
 ["22-25","jnz_tf_nor_yqdjb_1_ahqsgy","暗黑骑士盖亚","4%","全军破甲杀伤+5"],
 ["26-29","jnz_tf_nor_yqdjb_1_lphzjs","龙破坏之剑士","4%","全军破甲杀伤+5"],
 ["30-37","jnz_tf_nor_yqdjb_1_zzzl","诅咒之龙","8%","全军领导力+8"],
 ["38-45","jnz_tf_nor_yqdjb_1_hmdlh","黑魔导女孩","8%","所有法术冷却-10%"],
 ["46-50","jnz_tf_nor_yqdjb_1_youzu","被封印者的右足","5%","全军物理抗性+5%"],
 ["51-55","jnz_tf_nor_yqdjb_1_zuozu","被封印者的左足","5%","全军远程抗性+5%"],
 ["56-60","jnz_tf_nor_yqdjb_1_youwan","被封印者的右腕","5%","全军魔法抗性+5%"],
 ["61-65","jnz_tf_nor_yqdjb_1_zuowan","被封印者的左腕","5%","全军火焰抗性+5%"],
 ["66-70","jnz_tf_nor_yqdjb_1_tou","被封印的艾克佐迪亚","5%","全军伤害减免+5%；自身特殊保护-5"],
 ["71-78","jnz_tf_nor_yqdjb_1_ssfhz","神圣反护罩-反射镜力","8%","全军获得技能 jn_yqdjb_ssfhz"],
 ["79-83","jnz_tf_nor_yqdjb_1_ronghe","融合","5%","触发融合机制（自定义效果）"],
 ["84-91","jnz_tf_nor_yqdjb_1_heidong","黑洞","8%","全军魔法抗性+10%"],
 ["92-94","jnz_tf_nor_yqdjb_1_jsb","欧贝利斯克的巨神兵","3%","全军破甲杀伤+10，启用「谋杀冲锋」"],
 ["95-97","jnz_tf_nor_yqdjb_1_tkl","奥西里斯的天空龙","3%","全军破甲远程伤害+25%，射程+10%"],
 ["98-100","jnz_tf_nor_yqdjb_1_ysl","太阳神的翼神龙","3%","全军护甲+15，火焰抗性+20%"],
 ["组合","yqdjb_qnjm","艾克佐迪亚集齐","-","右足+左足+右腕+左腕+头 → 特殊奖励"],
 ["组合","yqdjb_lqsgy","龙骑士盖亚","-","真红骑士+暗黑骑士+融合 → 领袖武器强度+100，破甲+200"],
 ["组合","yqdjb_cmdjs","超魔导剑士-黑暗帕拉丁","-","黑魔术师+龙破坏剑士+融合"],
 ["组合","三幻神换身","巨神兵+天空龙+翼神龙","-","集齐三幻神 → 换身奖励"],
]
make_sheet("游戏王抽卡", ["随机区间/组合", "条目key", "显示名", "概率", "效果"], ygo,
           [14, 34, 26, 8, 70])

# =================================================================
# Sheet 6: 祈愿装备（唯一物品）
# =================================================================
anc_eff = {}
for j in read_db_dir("ancillary_to_effects_tables"):
    anc_eff.setdefault(j.get("ancillary", j.get("ancillary_key", "")), []).append(j)
wy_rows = []
WY_NOTES = {
 "ex_wy_xzzs": "档位1；觉醒版需帕秋莉", "ex_wy_xzzs_wjx": "档位1未觉醒版",
 "ex_wy_zaz": "档位2", "ex_wy_zaz_wjx": "档位2未觉醒",
 "ex_wy_sss": "档位3", "ex_wy_sss_wjx": "档位3未觉醒",
 "ex_wy_hhjj": "档位4", "ex_wy_hhjj_wjx": "档位4未觉醒",
 "ex_wy_xwzm": "档位5", "ex_wy_xwzm_wjx": "档位5未觉醒",
 "ex_wy_yfdxz": "档位6-7：全玩家发放",
 "ex_wy_zxqj": "档位8：全玩家发放，触发星の祈愿",
}
for k in ["ex_wy_xzzs","ex_wy_xzzs_wjx","ex_wy_zaz","ex_wy_zaz_wjx","ex_wy_sss",
          "ex_wy_sss_wjx","ex_wy_hhjj","ex_wy_hhjj_wjx","ex_wy_xwzm","ex_wy_xwzm_wjx",
          "ex_wy_yfdxz","ex_wy_zxqj"]:
    effs = []
    for j in anc_eff.get(k, []):
        ek = j.get("effect_key", "") or j.get("effect", "")
        scope = j.get("effect_scope", "")
        d = clean(loc.get("effects_description_" + ek, ""))
        v = j.get("value", "")
        head = f"{ek}" if not d else d
        effs.append(f"{head} = {v}（{scope}）" if v else head)
    wy_rows.append([k, name_of(k, "ancillaries_onscreen_name_"),
                    clean(loc.get("ancillaries_colour_text_" + k, "")),
                    WY_NOTES.get(k, ""), "\n".join(dict.fromkeys(effs)) or "见脚本效果"])
make_sheet("祈愿装备-唯一物品", ["物品key", "显示名", "物品描述", "抽取档位", "效果明细"], wy_rows,
           [22, 16, 40, 24, 70])

# =================================================================
# Sheet 7: 学生池（基沃托斯抽取）
# =================================================================
var_src = (LUA / "psw_变量增加.lua").read_text(encoding="utf-8")
stu_rows = []
for m in re.finditer(r'\["(psw_dlc7_jwts_mulu_\w+)"\]\s*=\s*\{(.*?)\n\s*\},', var_src, re.S):
    mulu, body = m.group(1), m.group(2)
    mulu_name = name_of(mulu, "rituals_display_name_", "") or mulu
    for cm_ in re.finditer(r'card_key\s*=\s*"([^"]+)"\s*,\s*types\s*=\s*"([^"]+)"\s*,\s*qj_num\s*=\s*(\d+)', body):
        ck, ty, qj = cm_.groups()
        stu_rows.append([mulu_name, ck, name_of(ck, "rituals_display_name_"),
                         ty, qj, clean(loc.get(ck + "_lore", ""))])
make_sheet("学生池-基沃托斯", ["社团", "学生key", "学生名", "类型", "奇迹计数", "效果说明"], stu_rows,
           [18, 40, 14, 8, 10, 70])

# =================================================================
wb.remove(wb.active)
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved:", OUT)
for ws in wb.worksheets:
    print(f"  {ws.title}: {ws.max_row - 1} rows")
